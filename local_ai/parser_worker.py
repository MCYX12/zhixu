"""Runs untrusted document parsing in a disposable, time-bounded subprocess."""

import io
import json
import sys
import zipfile

MAX_TEXT = 400_000


def parse(raw, suffix):
    pages = []
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ValueError("暂不支持加密 PDF，请先解密后上传")
        if len(reader.pages) > 150:
            raise ValueError("PDF 超过 150 页，请分卷上传")
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            pages.append({"title": f"第 {i + 1} 页", "text": text})
            if sum(len(p["text"]) for p in pages) > MAX_TEXT:
                raise ValueError("文档正文过长，请拆分后上传")
    elif suffix == ".docx":
        from docx import Document

        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 20_000_000:
                raise ValueError("DOCX 解压后过大，请拆分后上传")
        document = Document(io.BytesIO(raw))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        pages = [{"title": "", "text": "\n\n".join(parts)}]
    else:
        text = raw.decode("utf-8-sig")
        if "\x00" in text:
            raise ValueError("文件不是有效文本")
        pages = [{"title": "", "text": text}]
    if sum(len(p["text"]) for p in pages) > MAX_TEXT:
        raise ValueError("文档正文过长，请拆分后上传")
    if not any(p["text"].strip() for p in pages):
        raise ValueError("未提取到文字；扫描件请先进行 OCR")
    return pages


if __name__ == "__main__":
    try:
        import resource

        resource.setrlimit(resource.RLIMIT_CPU, (15, 15))
        raw = sys.stdin.buffer.read(5 * 1024 * 1024 + 1)
        if len(raw) > 5 * 1024 * 1024:
            raise ValueError("单个文档不能超过 5 MB")
        result = {"pages": parse(raw, sys.argv[1])}
    except ValueError as error:
        result = {"error": str(error)}
    except Exception:
        result = {"error": "无法解析该文件，请检查格式或导出为文本后重试"}
    print(json.dumps(result, ensure_ascii=False))
