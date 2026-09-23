import io
import json
import time

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from local_ai.api import create_app
from local_ai.config import Settings
from local_ai.workspace_store import WorkspaceStore


@pytest.fixture
def workspace(tmp_path):
    settings = Settings(database=tmp_path / "state.sqlite3", knowledge_root=tmp_path)
    app = create_app(
        settings,
        httpx.MockTransport(lambda r: httpx.Response(200, json={"data": [{"id": "qwen"}]})),
    )
    store = app.state.store
    keys = {
        name: store.create_user(
            name, ["lab"] if name != "outsider" else [], "admin" if name == "admin" else "member"
        )
        for name in ["admin", "alice", "bob", "outsider"]
    }
    with TestClient(app) as client:
        client.headers["Authorization"] = "Bearer " + keys["alice"]
        yield client, store, keys


def upload(
    client,
    filename="notes.md",
    content=b"# Lab\n\nProject ZEUS launch date is October 3.",
    visibility="private",
    workspace_id=None,
):
    params = {"filename": filename, "visibility": visibility}
    if workspace_id:
        params["workspace_id"] = workspace_id
    return client.post(
        "/api/documents",
        params=params,
        content=content,
        headers={"Content-Type": "application/octet-stream"},
    )


def test_cookie_session_csrf_origin_expiry_logout(workspace):
    client, store, keys = workspace
    result = client.post("/api/session")
    assert result.status_code == 200
    assert "HttpOnly" in result.headers["set-cookie"]
    csrf = result.json()["csrf"]
    del client.headers["Authorization"]
    assert client.get("/api/me").json()["user"]["id"] == "alice"
    assert client.delete("/api/session").status_code == 403
    assert (
        client.delete(
            "/api/session", headers={"X-CSRF-Token": csrf, "Origin": "https://evil.test"}
        ).status_code
        == 403
    )
    assert client.delete("/api/session", headers={"X-CSRF-Token": csrf}).status_code == 200
    assert client.get("/api/me").status_code == 401
    client.headers["Authorization"] = "Bearer " + keys["alice"]
    client.post("/api/session")
    del client.headers["Authorization"]
    with store.connection() as db:
        db.execute("UPDATE web_sessions SET expires_at=?", (time.time() - 1,))
    assert client.get("/api/me").status_code == 401


def test_disabled_account_revokes_cookie(workspace):
    client, store, keys = workspace
    client.post("/api/session")
    del client.headers["Authorization"]
    store.disable_user("alice")
    assert client.get("/api/me").status_code == 401


def test_upload_private_acl_preview_original_delete(workspace):
    client, store, keys = workspace
    r = upload(client)
    assert r.status_code == 201, r.text
    doc = r.json()
    did = doc["id"]
    assert doc["chunk_count"] == 1 and doc["visibility"] == "private"
    assert client.get("/api/documents/" + did).json()["chunks"][0]["text"].startswith("# Lab")
    assert b"ZEUS" in client.get("/api/documents/" + did + "/download").content
    client.headers["Authorization"] = "Bearer " + keys["bob"]
    assert client.get("/api/documents").json()["data"] == []
    assert client.get("/api/documents/" + did).status_code == 404
    assert client.get("/api/documents/" + did + "/download").status_code == 404
    assert client.delete("/api/documents/" + did).status_code == 404
    assert (
        client.patch(
            "/api/documents/" + did, json={"visibility": "group", "workspace_id": "lab"}
        ).status_code
        == 404
    )
    client.headers["Authorization"] = "Bearer " + keys["alice"]
    assert client.delete("/api/documents/" + did).status_code == 200
    assert client.get("/api/documents/" + did + "/download").status_code == 404
    with store.connection() as db:
        assert db.execute("SELECT COUNT(*) FROM document_files").fetchone()[0] == 0
        assert db.execute("SELECT COUNT(*) FROM chunks").fetchone()[0] == 0


def test_group_and_public_publication_controls(workspace):
    client, store, keys = workspace
    assert upload(client, visibility="public").status_code == 403
    assert upload(client, visibility="group", workspace_id="other").status_code == 403
    doc = upload(client, visibility="group", workspace_id="lab").json()
    client.headers["Authorization"] = "Bearer " + keys["bob"]
    assert client.get("/api/documents/" + doc["id"]).status_code == 200
    assert (
        client.patch("/api/documents/" + doc["id"], json={"visibility": "private"}).status_code
        == 403
    )
    client.headers["Authorization"] = "Bearer " + keys["outsider"]
    assert client.get("/api/documents/" + doc["id"]).status_code == 404
    client.headers["Authorization"] = "Bearer " + keys["admin"]
    assert upload(client, visibility="public").status_code == 201
    # Administrator role does not bypass private retrieval permissions.
    client.headers["Authorization"] = "Bearer " + keys["alice"]
    private = upload(client, filename="private.md").json()
    client.headers["Authorization"] = "Bearer " + keys["admin"]
    assert client.get("/api/documents/" + private["id"]).status_code == 404


def test_acl_revocation_removes_search_and_sources(workspace):
    client, store, keys = workspace
    doc = upload(client, visibility="group", workspace_id="lab").json()
    chunk = client.get("/api/documents/" + doc["id"]).json()["chunks"][0]["id"]
    client.headers["Authorization"] = "Bearer " + keys["bob"]
    assert client.get("/v1/sources/" + chunk).status_code == 200
    client.headers["Authorization"] = "Bearer " + keys["alice"]
    assert (
        client.patch("/api/documents/" + doc["id"], json={"visibility": "private"}).status_code
        == 200
    )
    client.headers["Authorization"] = "Bearer " + keys["bob"]
    assert client.get("/v1/sources/" + chunk).status_code == 404
    assert client.post("/v1/knowledge/search", json={"query": "ZEUS"}).json()["data"] == []


def test_file_constraints_and_failed_parse_do_not_publish(workspace):
    client, store, keys = workspace
    assert upload(client, filename="../secret.md").status_code == 400
    assert upload(client, filename="x.html").status_code == 422
    assert upload(client, content=b"").status_code == 422
    assert upload(client, content=b"\xff\x00").status_code == 422
    assert upload(client, content=b"x" * (5 * 1024 * 1024 + 1)).status_code == 413
    assert client.get("/api/documents").json()["data"] == []


def test_docx_tables_and_pdf_empty_handling(workspace):
    client, store, keys = workspace
    doc = Document()
    doc.add_paragraph("Quarterly review ZEUS")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Budget"
    table.cell(0, 1).text = "42"
    data = io.BytesIO()
    doc.save(data)
    r = upload(client, filename="review.docx", content=data.getvalue())
    assert r.status_code == 201, r.text
    chunks = client.get("/api/documents/" + r.json()["id"]).json()["chunks"]
    assert "Budget | 42" in chunks[0]["text"]
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    pdf = io.BytesIO()
    writer.write(pdf)
    response = upload(client, filename="scan.pdf", content=pdf.getvalue())
    assert response.status_code == 422 and "OCR" in response.text


def test_duplicate_filename_does_not_replace(workspace):
    client, store, keys = workspace
    a = upload(client).json()
    b = upload(client, content=b"Changed text").json()
    assert a["id"] != b["id"] and a["path"] != b["path"]
    assert len(client.get("/api/documents").json()["data"]) == 2


def test_text_pdf_upload_preserves_page_and_original(workspace):
    client, store, keys = workspace
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 250 Td (ZEUS approval code 8492) Tj ET")
    page[NameObject("/Contents")] = stream
    output = io.BytesIO()
    writer.write(output)
    original = output.getvalue()
    response = upload(client, filename="approval.pdf", content=original)
    assert response.status_code == 201, response.text
    did = response.json()["id"]
    chunks = client.get("/api/documents/" + did).json()["chunks"]
    assert chunks[0]["title"] == "第 1 页"
    assert "ZEUS approval code 8492" in chunks[0]["text"]
    assert client.get("/api/documents/" + did + "/download").content == original


def test_https_session_cookie_is_secure(workspace):
    client, store, keys = workspace
    with TestClient(client.app, base_url="https://testserver") as secure_client:
        response = secure_client.post(
            "/api/session", headers={"Authorization": "Bearer " + keys["alice"]}
        )
        assert response.status_code == 200
        assert "Secure" in response.headers["set-cookie"]
        assert secure_client.get("/api/me").status_code == 200


def conversation_body():
    return {
        "title": "Project review",
        "mode": "knowledge",
        "messages": [
            {"role": "user", "content": "What is ZEUS?"},
            {"role": "assistant", "content": "A project."},
        ],
    }


def test_conversation_isolation_revision_and_delete(workspace):
    client, store, keys = workspace
    r = client.post("/api/conversations", json=conversation_body())
    assert r.status_code == 201
    item = r.json()
    cid = item["id"]
    assert item["revision"] == 1
    client.headers["Authorization"] = "Bearer " + keys["bob"]
    assert client.get("/api/conversations").json()["data"] == []
    assert client.get("/api/conversations/" + cid).status_code == 404
    assert client.put("/api/conversations/" + cid, json=conversation_body()).status_code == 409
    assert client.delete("/api/conversations/" + cid).status_code == 404
    client.headers["Authorization"] = "Bearer " + keys["alice"]
    body = conversation_body() | {"revision": 1, "title": "Renamed"}
    assert client.put("/api/conversations/" + cid, json=body).json()["revision"] == 2
    assert client.put("/api/conversations/" + cid, json=body).status_code == 409
    assert client.delete("/api/conversations/" + cid).status_code == 200


def test_forged_sources_not_saved_and_revoked_sources_filtered(workspace):
    client, store, keys = workspace
    doc = upload(client, visibility="group", workspace_id="lab").json()
    chunk = client.get("/api/documents/" + doc["id"]).json()["chunks"][0]["id"]
    client.headers["Authorization"] = "Bearer " + keys["bob"]
    body = conversation_body()
    body["messages"][1]["sources"] = [
        {"chunk_id": chunk, "path": "forged", "url": "https://evil.test"},
        {"chunk_id": "nonexistent"},
    ]
    item = client.post("/api/conversations", json=body).json()
    assert len(item["messages"][1]["sources"]) == 1
    assert item["messages"][1]["sources"][0]["path"] == "notes.md"
    client.headers["Authorization"] = "Bearer " + keys["alice"]
    client.patch("/api/documents/" + doc["id"], json={"visibility": "private"})
    client.headers["Authorization"] = "Bearer " + keys["bob"]
    assert client.get("/api/conversations/" + item["id"]).json()["messages"][1]["sources"] == []


def test_static_assets_and_csp(workspace):
    client, store, keys = workspace
    response = client.get("/")
    assert response.status_code == 200
    assert "connect-src 'self'" in response.headers["content-security-policy"]
    for path in ["app.js", "app.css", "brand.svg", "vendor/purify.min.js", "vendor/marked.umd.js"]:
        assert client.get("/static/" + path).status_code == 200
    assert "https://cdn" not in response.text


def test_audit_no_document_content_and_backup(workspace, tmp_path):
    client, store, keys = workspace
    upload(client, content=b"ULTRA_SECRET_PAYLOAD")
    status = client.get("/api/status").json()
    assert status["model_ready"]
    assert "ULTRA_SECRET_PAYLOAD" not in json.dumps(status)
    assert status["events"][0]["action"] == "document.upload"
    target = tmp_path / "backup.sqlite3"
    store.backup(target)
    restored = WorkspaceStore(target)
    alice = restored.authenticate(keys["alice"])
    assert restored.documents(alice)[0]["chunk_count"] == 1


def test_backup_status_visible_only_to_admin(workspace, monkeypatch):
    client, store, keys = workspace
    monkeypatch.setattr(
        "local_ai.workspace_api.automatic_backup_status",
        lambda: {"state": "ok", "last_success": 123, "last_attempt": 123},
    )
    assert client.get("/api/status").json()["backup"] is None
    client.headers["Authorization"] = "Bearer " + keys["admin"]
    assert client.get("/api/status").json()["backup"]["state"] == "ok"


def test_member_management_admin_only_and_immediate_revocation(workspace):
    client, store, keys = workspace
    change = {"role": "member", "groups": [], "active": False}
    assert client.get("/api/members").status_code == 403
    assert client.patch("/api/members/bob", json=change).status_code == 403
    bob = store.authenticate(keys["bob"])
    token, _ = store.create_session(bob)
    client.headers["Authorization"] = "Bearer " + keys["admin"]
    result = client.get("/api/members")
    assert result.status_code == 200
    assert all(set(m) == {"id", "role", "active", "groups"} for m in result.json()["data"])
    assert client.patch("/api/members/bob", json=change).status_code == 200
    assert store.authenticate(keys["bob"]) is None and store.session(token) is None
    change["active"] = True
    assert client.patch("/api/members/bob", json=change).status_code == 200
    assert store.authenticate(keys["bob"]).groups == ()
    assert store.session(token) is None
    assert client.patch("/api/members/admin", json=change).status_code == 409
    assert client.patch("/api/members/missing", json=change).status_code == 404
    assert client.patch("/api/members/bob", json=change | {"groups": ["../bad"]}).status_code == 400


def test_member_updates_require_csrf_with_cookie(workspace):
    client, store, keys = workspace
    client.headers["Authorization"] = "Bearer " + keys["admin"]
    csrf = client.post("/api/session").json()["csrf"]
    del client.headers["Authorization"]
    change = {"role": "member", "active": True, "groups": ["new_group"]}
    assert client.patch("/api/members/bob", json=change).status_code == 403
    assert (
        client.patch("/api/members/bob", json=change, headers={"X-CSRF-Token": csrf}).status_code
        == 200
    )


def test_group_removal_revokes_document_and_source_access(workspace):
    client, store, keys = workspace
    response = upload(client, visibility="group", workspace_id="lab")
    assert response.status_code == 201
    document_id = response.json()["id"]
    client.headers["Authorization"] = "Bearer " + keys["bob"]
    assert client.get("/api/documents/" + document_id).status_code == 200
    client.headers["Authorization"] = "Bearer " + keys["admin"]
    assert (
        client.patch(
            "/api/members/bob", json={"role": "member", "groups": [], "active": True}
        ).status_code
        == 200
    )
    client.headers["Authorization"] = "Bearer " + keys["bob"]
    assert client.get("/api/documents/" + document_id).status_code == 404
    assert client.get("/api/documents/" + document_id + "/download").status_code == 404
